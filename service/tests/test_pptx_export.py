"""The outbound pptx bridge: what reaches the slide, and what says it didn't.

Silent loss is the worst failure this module can have, because the exported
deck still opens, still looks plausible, and the missing paragraph is only
discovered by the person presenting it. It is the failure IMPORT.md rules out
("Never: a plausible-looking reconstruction that silently drops what it didn't
understand"), and the exporter used to have it ten different ways at once.

Two invariants stand behind these cases:

  coverage — `_partition` maps every element of a slide to a part, so
    `unplaced_text()` is empty for any slide, whatever classes it carries.
    The profile permits unknown classes by design, so an exporter that draws
    only what a fixed list of `find()` calls names is unsound by
    construction; the mapping has to be total.

  tiling — `_stack` deals bands down a column off one cursor, so two of them
    cannot overlap and none can leave the column. A short measurement makes a
    band tight, never makes it land on its neighbour.

These need no model, no network and no Drive: a dialect string goes in, a pptx
comes back, and the assertions read the shapes. python-pptx is a hard
dependency of the service, so there is nothing to skip.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from dia_service.pptx_export import (
    ADVANCE_EM,
    ADVANCE_EM_CAPS,
    DESIGN_H,
    DESIGN_W,
    Band,
    _parse,
    _read_theme,
    _stack,
    _table_plan,
    _text_h,
    deck_to_pptx,
    unplaced_text,
)

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"

THEME = (
    '<style id="dia-theme">:root{'
    "--dia-paper:#FFFFFF;--dia-ink:#202124;--dia-ink-soft:#3C4043;"
    "--dia-ink-faint:#5F6368;--dia-accent:#4285F4;--dia-rule:#DADCE0;"
    "--dia-scale-1:15px;--dia-scale-2:16px;--dia-scale-4:20px;"
    "--dia-scale-5:34px;--dia-gap:24px;--dia-pad:53px;}</style>"
)


def deck(body: str) -> str:
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<title>t</title>{THEME}</head><body>"
        f'<section class="dia-slide">{body}</section>'
        "</body></html>"
    )


def render(body: str):
    """(presentation, warnings) for a one-slide dialect deck."""
    warnings: list[str] = []
    prs = Presentation(io.BytesIO(deck_to_pptx(deck(body), warnings)))
    return prs, warnings


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6 and hasattr(sh, "shapes"):  # GROUP
            yield from walk(sh.shapes)


def all_text(prs) -> str:
    """Every string that reaches the canvas. Speaker notes are deliberately
    excluded: text parked in the notes is not text the audience can read."""
    out = []
    for slide in prs.slides:
        for sh in walk(slide.shapes):
            if sh.has_table:
                out += [c.text for r in sh.table.rows for c in r.cells]
            elif sh.has_text_frame:
                out.append(sh.text_frame.text)
    return " ".join(out)


def rect(sh):
    return (
        sh.left / 9525,
        sh.top / 9525,
        (sh.width or 0) / 9525,
        (sh.height or 0) / 9525,
    )


TITLE = '<h2 class="dia-title">A title</h2>'


# ---------------------------------------------------------------------------
# coverage: nothing on a slide may go missing without a word
# ---------------------------------------------------------------------------

# One canary per region that used to leave the deck in silence. Every entry
# here was measured dropping on the shape of this module before the partition
# replaced the role enumeration -- ten roles, one bug.
DROPPED = [
    ("subtitle", '<p class="dia-subtitle">CANARY</p>'),
    ("quote", '<blockquote class="dia-quote">CANARY</blockquote>'),
    ("attribution", '<p class="dia-attribution">CANARY</p>'),
    ("code block", '<pre class="dia-code">CANARY</pre>'),
    ("island", '<div class="dia-island"><p>CANARY</p></div>'),
    ("list outside a body", '<ul class="dia-list"><li>CANARY</li></ul>'),
    (
        "second table",
        "<table><tr><td>first</td></tr></table>"
        "<table><tr><td>CANARY</td></tr></table>",
    ),
    (
        "second figure",
        '<figure class="dia-figure"><img src="a" alt="one"></figure>'
        '<figure class="dia-figure"><img src="b" alt="CANARY"></figure>',
    ),
    ("bare paragraph", "<p>CANARY</p>"),
    (
        "second column",
        '<div class="dia-columns"><div class="dia-col"><p>first</p></div>'
        '<div class="dia-col"><p>CANARY</p></div></div>',
    ),
    ("unknown class", '<div class="marketing-callout">CANARY</div>'),
]


@pytest.mark.parametrize("name,frag", DROPPED, ids=[n for n, _ in DROPPED])
def test_a_region_never_leaves_the_deck_in_silence(name, frag):
    """Every one of these used to vanish: no error, no warning, no marker.

    The island case is the sharpest, because `translate-slide` rule 5 tells
    the model to island whatever it cannot express -- so the more honest the
    translation, the more of the deck went missing."""
    prs, _ = render(TITLE + frag)
    assert "CANARY" in all_text(prs), f"{name} did not reach the slide"


def test_the_partition_accounts_for_every_word():
    """The invariant the coverage cases rest on, read directly.

    `unplaced_text` is empty for any slide -- including one built entirely
    from classes and tags this module has never heard of, which the profile
    explicitly allows ("unknown classes are permitted ... and are not
    flagged")."""
    root = _parse(
        deck(
            '<div class="brandwrap"><section class="teaser">alpha'
            "<span>bravo</span></section>"
            '<article data-role="x"><p>charlie</p>loose delta</article></div>'
        )
    )
    slide = root.find(lambda e: e.has("dia-slide"))
    assert unplaced_text(slide) == ""


def test_a_style_block_is_accounted_for_without_being_printed():
    """`mute` is the only thing the partition drops, and it drops it on
    purpose. The distinction matters: an omission looks exactly like this
    until the day it eats a paragraph."""
    body = "<style>.x{color:red}</style><p>visible copy</p>"
    root = _parse(deck(body))
    slide = root.find(lambda e: e.has("dia-slide"))
    assert unplaced_text(slide) == ""
    prs, _ = render(body)
    text = all_text(prs)
    assert "visible copy" in text
    assert "color:red" not in text


def test_mixed_inline_content_sets_as_one_paragraph():
    """A wrapper whose descendants carry no band of their own is one run of
    prose. Descending into it would set the text either side of an `<em>` as
    separate paragraphs."""
    prs, _ = render('<div class="note">then <em>measures</em> everything.</div>')
    frames = [
        sh.text_frame for sh in walk(prs.slides[0].shapes)
        if sh.has_text_frame and "measures" in sh.text_frame.text
    ]
    assert len(frames) == 1
    assert len(frames[0].paragraphs) == 1
    assert frames[0].text.strip() == "then measures everything."


def test_a_recovered_island_says_so():
    """An island is markup the translation could not express, so its text is
    on the slide but its appearance is not. That is a real loss of fidelity
    and the sink has to name it -- unlike a second table or a second column,
    which now render natively and need no apology."""
    _, warnings = render(TITLE + '<div class="dia-island"><p>x</p></div>')
    assert warnings and "island" in warnings[0]


def test_a_second_table_renders_as_a_table():
    """`_figure_of()` returned the FIRST table/img/svg and every later one was
    dropped. Recovering its text would be enough to stop the loss; drawing it
    as a native table is what the format can actually carry, so the warning
    sink stays quiet."""
    prs, warnings = render(
        TITLE
        + "<table><tr><td>first_a</td></tr></table>"
        + "<table><tr><td>second_b</td></tr></table>"
    )
    tables = [sh for sh in walk(prs.slides[0].shapes) if sh.has_table]
    assert len(tables) == 2
    text = all_text(prs)
    assert "first_a" in text and "second_b" in text
    assert warnings == []


def test_subtitle_is_rendered():
    """`dia-subtitle` is in the vocabulary translate-slide tells the model to
    emit, but `_render_slide` never looked for it, so it was always dropped.
    This is the one guaranteed to bite anyone following the skill as written."""
    prs, _ = render(TITLE + '<p class="dia-subtitle">SUBCANARY text</p>')
    assert "SUBCANARY" in all_text(prs)


@pytest.mark.parametrize("name", ["what-is-dia.html", "demo-deck.html"])
def test_an_example_deck_loses_no_prose(name):
    """End to end on the decks the repository ships, which is where the size
    of the old failure shows: `what-is-dia.html` lost 403 distinct words on
    its way to a .pptx, and the warning sink said nothing about any of them.

    The comparison is a word of the source against the LETTERS of the export,
    not word against word: adjacent inline elements ("<span>research</span>
    <span>preview</span>") read back from the source as one run even though
    nothing was lost. Text inside an `<svg>` is left out for the other half of
    the same problem -- a scene draws each `<text>` node as its own label, so
    the pieces reach the slide in boxes rather than in one stream. Loss is
    what this measures, not tokenization."""
    html = (EXAMPLES / name).read_text(encoding="utf-8")
    root = _parse(html)
    slides = root.find_all(lambda e: e.tag == "section" and e.has("dia-slide"))
    prs = Presentation(io.BytesIO(deck_to_pptx(html)))

    for i, (src_el, out) in enumerate(zip(slides, prs.slides), 1):
        got = []
        for sh in walk(out.shapes):
            if sh.has_table:
                got += [c.text for r in sh.table.rows for c in r.cells]
            elif sh.has_text_frame:
                got.append(sh.text_frame.text)
        got.append(out.notes_slide.notes_text_frame.text)
        letters = re.sub(r"[^a-z]+", "", " ".join(got).lower())
        src = " ".join(
            e.text for e in src_el.walk() if e.text and not _inside_svg(e)
        )
        missing = [
            word
            for word in set(re.findall(r"[a-z]{4,}", src.lower()))
            if word not in letters
        ]
        assert not missing, f"slide {i} lost {sorted(missing)}"


def _inside_svg(el) -> bool:
    p = el.parent
    while p is not None:
        if p.tag == "svg":
            return True
        p = p.parent
    return False


# ---------------------------------------------------------------------------
# tiling: bands are disjoint, and inside the space they were given
# ---------------------------------------------------------------------------


def _probe_bands(asks, floors=None, grow=(), push=()):
    placed = []

    def make(i, ask, floor):
        def paint(_slide, _theme, x, y, w, h):
            placed.append((x, y, w, h))

        return Band(
            f"b{i}", ask, paint, floor=floor, grow=i in grow, push=i in push,
            gap_after=8.0,
        )

    floors = floors or [min(a, 20.0) for a in asks]
    bands = [make(i, a, f) for i, (a, f) in enumerate(zip(asks, floors))]
    return bands, placed


@pytest.mark.parametrize(
    "asks",
    [
        [40, 60, 30],
        [400, 400, 400],  # over-subscribed: must compress, not overlap
        [10],
        [5000, 10],  # absurdly over-subscribed
        [30, 30, 30, 30, 30, 30, 30, 30],
    ],
)
def test_bands_tile_their_column(asks):
    """Rects come off one cursor, so they are ordered, disjoint and inside the
    column whatever the measurement said. The old layout advanced a cursor by
    a character-count estimate and drew wherever it landed, so a band whose
    estimate was short was drawn under by the next one."""
    bands, placed = _probe_bands(asks)
    _stack(None, None, bands, 50.0, 60.0, 600.0, 500.0)
    assert len(placed) == len(asks)
    for (_, y, _, h) in placed:
        assert y >= 60.0 - 0.01
        assert y + h <= 60.0 + 500.0 + 0.01
    for a, b in zip(placed, placed[1:]):
        assert a[1] + a[3] <= b[1] + 0.01, "two bands overlap"


def test_an_over_subscribed_column_reports_the_shortfall():
    """Compressing bands to keep them on the page is a visible compromise, so
    it has to be said out loud rather than looking deliberate."""
    bands, _ = _probe_bands([600, 600, 600], floors=[300, 300, 300])
    assert _stack(None, None, bands, 0.0, 0.0, 500.0, 400.0) > 1.0
    bands, _ = _probe_bands([40, 60])
    assert _stack(None, None, bands, 0.0, 0.0, 500.0, 400.0) == 0.0


def test_a_trailing_caption_sits_at_the_foot_but_not_when_centred():
    """`.foot { margin-top: auto }` is the dialect's own idiom, and a centred
    column already has a placement rule -- letting both spend the same slack
    put a cover's caption 90px below the bottom of the page."""
    bands, placed = _probe_bands([40, 20], push={1})
    _stack(None, None, bands, 0.0, 0.0, 600.0, 400.0)
    assert placed[1][1] + placed[1][3] == pytest.approx(400.0)
    bands, placed = _probe_bands([40, 20], push={1})
    _stack(None, None, bands, 0.0, 0.0, 600.0, 400.0, center=True)
    assert placed[-1][1] + placed[-1][3] <= 400.0


# ---------------------------------------------------------------------------
# the table, and the collision that reserving a band underneath it used to make
# ---------------------------------------------------------------------------


def _theme():
    return _read_theme(_parse(deck("")))


def test_a_table_occupies_the_height_the_layout_reserved():
    """A table's height is not a guess made at the call site: the rows the
    plan measures are written back as explicit row heights, so the shape is
    the size the layout set aside for it. Leaving the rows to the renderer is
    what let a table grow past its box."""
    tbl_html = "<table>" + "".join(
        f"<tr><td>row {i} with enough text in it to wrap once or twice at "
        f"this width</td><td>{i}</td></tr>"
        for i in range(6)
    ) + "</table>"
    root = _parse(deck(tbl_html))
    tbl = root.find(lambda e: e.tag == "table")
    plan = _table_plan(_theme(), tbl, 600.0)
    assert plan.height == pytest.approx(sum(plan.rows))

    prs, _ = render(TITLE + tbl_html)
    shape = next(sh for sh in walk(prs.slides[0].shapes) if sh.has_table)
    rows = sum(r.height for r in shape.table.rows) / 9525
    assert rows == pytest.approx(shape.height / 9525, abs=1.0)


def test_a_table_row_is_tall_enough_for_its_own_text():
    """The reservation is only worth anything if the rows really hold their
    cells. A flat per-row height -- what the old code assumed when it sized a
    table at 34px a row -- is short for any cell that wraps, and a renderer
    that grows the row to fit puts the table past the box the layout set
    aside for it."""
    theme = _theme()
    long_cell = (
        "a cell whose text is long enough that it has to wrap several times "
        "at any sensible column width, which is what a real table cell does"
    )
    root = _parse(deck(f"<table><tr><td>short</td></tr>"
                       f"<tr><td>{long_cell}</td></tr></table>"))
    tbl = root.find(lambda e: e.tag == "table")
    plan = _table_plan(theme, tbl, 400.0)
    needed = _text_h(long_cell, 400.0 - 16.0, theme.scale[2], 1.25)
    assert plan.rows[1] >= needed
    assert plan.rows[1] > plan.rows[0], "a wrapping row must be the taller one"


def test_a_table_too_tall_for_its_room_shrinks_and_says_so():
    """The alternative -- drawing it at full size over whatever is beneath --
    is the collision this measurement exists to prevent."""
    theme = _theme()
    root = _parse(
        deck("<table>" + "<tr><td>a cell with some text</td></tr>" * 40
             + "</table>")
    )
    tbl = root.find(lambda e: e.tag == "table")
    loose = _table_plan(theme, tbl, 600.0)
    tight = _table_plan(theme, tbl, 600.0, 200.0)
    assert loose.fits and not tight.fits
    assert tight.height < loose.height
    assert tight.scale < 1.0


def test_a_long_footnote_does_not_collide_with_a_big_table():
    """The limitation the fix for the footnote band originally left behind:
    the exporter estimated a table's height from character counts while the
    renderer laid the real table out taller, so reserving the footnote band
    converted clipping into collision.

    What holds this now is structure rather than a better estimate. The
    footnote's room comes out of the pool before anything else is measured,
    and the content bands come off one cursor inside what is left, so no
    measurement can put them in the same place. The measurement's job is the
    separate one two tests up: making the rows tall enough that the renderer
    does not grow the table past the box it was given."""
    prs, _ = render(
        TITLE
        + "<table>"
        + "".join(
            f"<tr><td>row {i} carries a sentence long enough to wrap</td>"
            f"<td>{i * 37}</td></tr>"
            for i in range(12)
        )
        + "</table>"
        + '<p class="dia-footnote">' + ("Caveat words follow. " * 40) + "</p>"
    )
    shapes = list(walk(prs.slides[0].shapes))
    table = next(sh for sh in shapes if sh.has_table)
    foot = next(
        sh for sh in shapes
        if sh.has_text_frame and "Caveat" in sh.text_frame.text
    )
    t = rect(table)
    f = rect(foot)
    assert t[1] + t[3] <= f[1] + 0.5, (
        f"the table runs to {t[1] + t[3]:.0f}px and the footnote starts at "
        f"{f[1]:.0f}px"
    )


# ---------------------------------------------------------------------------
# the footnote band
# ---------------------------------------------------------------------------


def footnote_box(prs):
    for sh in walk(prs.slides[0].shapes):
        if sh.has_text_frame and "Provenance" in sh.text_frame.text:
            return sh
    pytest.fail("the footnote never reached the slide at all")


LONG_NOTE = "Provenance " + (
    "sentence that keeps going and carries the caveats. " * 18
)


def test_footnote_box_grows_with_its_text():
    """The footnote used to be drawn in a FIXED 20px box pinned at
    `DESIGN_H - pad + 4`, whatever it contained. Real provenance notes run to
    ~1000 characters, so the tail rendered outside its own box and the
    renderer clipped it -- the study's caveats were the part that vanished.

    Note the shape of this assertion. Checking that the BOX sits inside the
    page does NOT catch the bug: a 20px box at y=672 is comfortably on the
    page, and the old code passed that check while losing most of the text.
    What distinguishes fixed-height from sized is whether the box responds to
    how much text it holds at all."""
    short = render(TITLE + '<p class="dia-footnote">Provenance, briefly.</p>')[0]
    long_ = render(TITLE + f'<p class="dia-footnote">{LONG_NOTE}</p>')[0]
    h_short = int(footnote_box(short).height)
    h_long = int(footnote_box(long_).height)
    assert h_long > h_short * 3, (
        "the footnote box does not grow with its text, so a long note "
        f"overflows it: short={Emu(h_short).pt:.0f}pt long={Emu(h_long).pt:.0f}pt"
    )


def test_footnote_stays_inside_the_page():
    """Sizing the box must not push it off the bottom instead."""
    prs, _ = render(TITLE + f'<p class="dia-footnote">{LONG_NOTE}</p>')
    sh = footnote_box(prs)
    page_h = int(Emu(DESIGN_H * 9525))
    assert int(sh.top) >= 0
    assert int(sh.top) + int(sh.height) <= page_h, (
        "footnote runs past the bottom edge: "
        f"bottom={Emu(int(sh.top) + int(sh.height)).pt:.0f}pt "
        f"page={Emu(page_h).pt:.0f}pt"
    )


def test_a_runaway_footnote_shrinks_rather_than_eating_the_slide():
    """A third of the slide is as much as a footnote may take before it IS
    the slide."""
    prs, _ = render(TITLE + '<p class="dia-footnote">Provenance '
                    + ("and one more caveat after another. " * 200) + "</p>")
    assert footnote_box(prs).height / 9525 <= DESIGN_H * 0.35


# ---------------------------------------------------------------------------
# the page as a frame
# ---------------------------------------------------------------------------


def test_everything_lands_inside_the_frame():
    """A dense slide -- subtitle, table, island and a long footnote together --
    must not put any shape outside the page."""
    prs, _ = render(
        TITLE
        + '<p class="dia-subtitle">Subtitle line</p>'
        + '<div class="dia-body"><p>Body copy.</p></div>'
        + "<table><tr><th>H</th></tr><tr><td>cell</td></tr></table>"
        + '<div class="dia-island"><p>island text</p></div>'
        + '<p class="dia-footnote">' + ("Caveat words. " * 30) + "</p>"
    )
    for sh in walk(prs.slides[0].shapes):
        if not (sh.width and sh.height):
            continue
        x, y, w, h = rect(sh)
        assert x >= -0.5 and y >= -0.5
        assert x + w <= DESIGN_W + 0.5
        assert y + h <= DESIGN_H + 0.5


@pytest.mark.parametrize("name", ["what-is-dia.html", "demo-deck.html"])
def test_an_example_deck_puts_nothing_off_the_page(name):
    """Including scene labels: an svg `<text>` used to get a fixed 420px box
    hung off the glyph origin, which ran off the page for anything in the
    right half of a diagram."""
    html = (EXAMPLES / name).read_text(encoding="utf-8")
    prs = Presentation(io.BytesIO(deck_to_pptx(html)))
    for i, slide in enumerate(prs.slides, 1):
        for sh in walk(slide.shapes):
            if not (sh.width and sh.height):
                continue
            x, y, w, h = rect(sh)
            assert (
                x >= -0.5 and y >= -0.5
                and x + w <= DESIGN_W + 0.5 and y + h <= DESIGN_H + 0.5
            ), f"slide {i}: {x:.0f},{y:.0f} {w:.0f}x{h:.0f} leaves the page"


# A 40x10 red PNG: a box that is not 4:1 must not make it its own shape.
WIDE_PNG = (
    "iVBORw0KGgoAAAANSUhEUgAAACgAAAAKCAIAAABJ+IsHAAAAGklEQVR4nGP4z8Aw"
    "IGhgbB21eNTiUYuHhcUAaRuOgDO6NfgAAAAASUVORK5CYII="
)


def test_an_image_keeps_its_proportions():
    """The figure box is the room the layout had, not the picture's shape.
    Handing python-pptx both a width and a height stretches the image to
    fill it -- a 4:1 diagram came out square in a square box."""
    prs, _ = render(
        TITLE
        + '<figure class="dia-figure">'
        + f'<img src="data:image/png;base64,{WIDE_PNG}" alt="wide">'
        + "</figure>"
    )
    pics = [sh for sh in walk(prs.slides[0].shapes) if sh.shape_type == 13]
    assert len(pics) == 1
    _, _, w, h = rect(pics[0])
    assert w / h == pytest.approx(4.0, rel=0.02)


def test_an_svg_label_lands_where_its_anchor_puts_it():
    """`text-anchor` decides which edge of the label sits at x. Ignoring it
    drew every `middle`- and `end`-anchored label a box-width to the right of
    where the svg drew it -- 46 such labels across the two example decks."""
    svg = (
        '<figure class="dia-figure"><svg viewBox="0 0 400 200">'
        '<text x="200" y="100" text-anchor="middle">CENTRED</text>'
        '<text x="200" y="150" text-anchor="end">TRAILING</text>'
        '<text x="200" y="50">LEADING</text>'
        "</svg></figure>"
    )
    prs, _ = render(TITLE + svg)
    boxes = {
        sh.text_frame.text: rect(sh)
        for sh in walk(prs.slides[0].shapes)
        if sh.has_text_frame and sh.text_frame.text in
        ("CENTRED", "TRAILING", "LEADING")
    }
    assert len(boxes) == 3
    centre = boxes["CENTRED"][0] + boxes["CENTRED"][2] / 2
    trailing_end = boxes["TRAILING"][0] + boxes["TRAILING"][2]
    leading_start = boxes["LEADING"][0]
    assert abs(centre - trailing_end) < 10
    assert leading_start < centre < trailing_end + 10


# ---------------------------------------------------------------------------
# the measurement the whole layout rests on
# ---------------------------------------------------------------------------

# Average glyph advance in em/char, measured at 100px with PIL over the fonts
# a Linux box has by default. Liberation Sans carries Arial's metrics (this
# module's default face) and Liberation Mono carries Courier's; DejaVu is the
# usual fallback. The layout's constants have to sit ABOVE every one of these,
# because a short measurement is what puts one band's text on the next.
MEASURED_ADVANCE = {
    "lowercase prose, Liberation Sans": 0.450,
    "mixed prose, Liberation Sans": 0.411,
    "title case, Liberation Sans": 0.444,
    "digits, Liberation Sans": 0.482,
    "lowercase prose, Liberation Sans Bold": 0.487,
    "any text, Liberation Mono": 0.600,
    "lowercase prose, DejaVu Sans": 0.513,
    "digits, DejaVu Sans": 0.548,
}
MEASURED_ADVANCE_CAPS = {
    "all caps, Liberation Sans": 0.664,
    "all caps, Liberation Sans Bold": 0.670,
    "all caps, DejaVu Sans": 0.674,
}


def test_the_measure_errs_toward_whitespace():
    """Over-estimating a line leaves a gap; under-estimating puts text on top
    of the next band. The constants are set above the widest realistic case
    rather than at the average, and this is the record of what "widest" meant.

    The old estimates -- 0.52 em for body text, 0.50 for the footnote -- sat
    UNDER the monospace face the repository's own `what-is-dia` theme sets for
    prose, and 25% under the all-caps text this module produces itself."""
    assert ADVANCE_EM > max(MEASURED_ADVANCE.values())
    assert ADVANCE_EM_CAPS > max(MEASURED_ADVANCE_CAPS.values())


# ---------------------------------------------------------------------------
# the warning sink
# ---------------------------------------------------------------------------


def test_a_clean_slide_warns_about_nothing():
    """The warning sink must stay quiet on a slide the exporter placed fully,
    or callers will learn to ignore it."""
    _, warnings = render(TITLE + '<div class="dia-body"><p>Just prose.</p></div>')
    assert warnings == []


@pytest.mark.parametrize("name", ["what-is-dia.html", "demo-deck.html"])
def test_an_example_deck_warns_about_nothing(name):
    """The decks the repository ships export cleanly, so a warning from one
    of them means the exporter changed, not the deck."""
    warnings: list[str] = []
    deck_to_pptx((EXAMPLES / name).read_text(encoding="utf-8"), warnings)
    assert warnings == []


def test_warnings_are_optional():
    """The sink is opt-in; existing callers that pass nothing keep working."""
    assert deck_to_pptx(deck(TITLE)), "no-warnings call must still return bytes"


def test_the_download_endpoint_carries_what_it_could_not_place():
    """A .pptx response has no body left to say it in, so the count and the
    first note ride on headers. A sink no caller reads is the silent failure
    with extra steps."""
    pytest.importorskip("httpx", reason="TestClient needs httpx")
    from fastapi.testclient import TestClient

    from dia_service import main

    body = deck(TITLE + '<div class="dia-island"><p>a widget</p></div>')
    with TestClient(main.app, base_url="http://127.0.0.1:8317") as c:
        r = c.post("/export/pptx", json={"html": body})
    assert r.status_code == 200
    assert r.headers["x-dia-export-warnings"] == "1"
    assert "island" in r.headers["x-dia-export-warning-1"]
